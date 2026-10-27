import re, argparse
from pathlib import Path
import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def clean(v):
    return ILLEGAL.sub('', v) if isinstance(v, str) else v


def uniq(seq):
    seen, out = set(), []
    for s in seq:
        k = s.upper()
        if k not in seen:
            seen.add(k)
            out.append(s)
    return out


NUM = re.compile(r'^\d{1,3}(\.\d{1,3})*$')
ANCHOR = re.compile(r'^\s*(GENERAL\s+)?(NOTES?|REMARKS?|הערות|הערה)\s*[:：]?\s*$', re.I)

CATS = [
    ("Material",   r"MATERIAL|ALLOY|ALUM\b|ALUMIN|STEEL|TITAN|BRASS|COPPER|6061|7075|304|316|17-4|ASTM B|AMS \d"),
    ("Cleaning",   r"CLEAN|DEGREAS|ALKALINE|RINSE"),
    ("Coating",    r"COAT|ANODI|PLATING|PLATED|PASSIV|ALODINE|CHEM FILM|PAINT|PRIMER|OXIDE|SURFACE TREATMENT"),
    ("Heat treat", r"HEAT TREAT|HARDEN|HRC|TEMPER|ANNEAL|STRESS RELIEF|THERMAL"),
    ("Marking",    r"ENGRAV|LASER MARK|ETCH|STAMP|MARKING|MARK PART|IDENTIFICATION|TRACEABIL|MIL-STD-130"),
    ("Masking",    r"MASK|DO NOT (COAT|PLATE)|FREE OF (PAINT|COAT|PLATING)"),
    ("Machining",  r"MACHIN|BURR|RADII|RADIUS|DEBURR|BREAK.*EDGE"),
    ("Assembly",   r"INSTALL|INSERT|HELICOIL"),
    ("Inspection", r"INSPECT|CMM|NDT|X-RAY|FIRST ARTICLE|GAUGE|SAMPLING|QUALITY ASSURANCE|AQL"),
    ("Packing",    r"PACK|SHIPMENT|PRESERV"),
]


def classify(*texts):
    t = " ".join(texts).upper()
    return [c for c, rx in CATS if re.search(rx, t)]


# ---------- standards ----------
STD_RE = re.compile(
    r'\b(?:SAE-)?(?:AMS|FED)-STD-\d+[A-Z]?'
    r'|\b(?:SAE-)?AMS[- ]?(?:QQ-[A-Z]-)?\d+(?:/\d+)?[A-Z]?'
    r'|\bASTM[- ]?[A-Z]\s?\d+[A-Z]?'
    r'|\bMIL-[A-Z]+-\d+[A-Z]?'
    r'|\bQQ-[A-Z]-\d+(?:/\d+)?'
    r'|\b(?:EN|DIN|ISO|JIS|GOST)[- ]?\d+[\d\-]*', re.I)
PER_RE = re.compile(r'\bPER\s+([A-Z0-9][A-Z0-9./\-]*)', re.I)   # customer specs like PS-111.21


def get_standards(t):
    found = [m.group(0).strip() for m in STD_RE.finditer(t)]
    for m in PER_RE.finditer(t):
        tok = m.group(1).rstrip('.,;:)')
        if any(c.isdigit() for c in tok):
            found.append(tok)
    found = uniq(found)
    return [s for s in found if not any(s != o and s.upper() in o.upper() for o in found)]


# ---------- material field extraction ----------
AL_RE = re.compile(r'\b([1-8]\d{3})(?:\s*-?\s*(T\d{1,5}|O|H\d{2,3}))?\b')
STEEL_RE = re.compile(r'\b(304L?|316L?|321|303|410|420|430|17-4\s?PH|15-5\s?PH|13-8\s?MO|'
                      r'4130|4140|4340|1018|1045|8620)\b', re.I)
TI_RE = re.compile(r'\bTI-?6AL-?4V\b|\bGR(?:ADE)?\.?\s?\d+\b', re.I)
TEMPER_RE = re.compile(r'\b(H\d{3,4}|ANNEALED|SOLUTION TREATED)\b', re.I)
FORM_RE = re.compile(r'\b(PLATE|SHEET|ROD|BAR|TUBE|TUBING|PIPE|ANGLE|CHANNEL|EXTRUSION|'
                     r'FORGING|CASTING|STRIP|WIRE|BILLET|BLOCK)\b', re.I)
DIM_RE = re.compile(r'\b(L|W|T|H|D|DIA|OD|ID)\s*=\s*([\d.]+)', re.I)
UNIT_RE = re.compile(r'\b(INCH(?:ES)?|MM)\b|\bIN\b', re.I)
CATNO_RE = re.compile(r'(\d{6,})\s+([\d.]+)\s*([A-Za-z]{1,4})?')


def extract_material(rows):
    mrows = [r for r in rows if "Material" in r["cat"]]
    if not mrows:
        return None
    text = " ".join(r["text"] for r in mrows)
    extra = " ".join(r["extra"] for r in mrows)
    stds = []
    for m in STD_RE.finditer(text):
        s = m.group(0).strip()
        if s not in stds:
            stds.append(s)
    body = STD_RE.sub(" ", text)
    up = body.upper()

    family, alloy, temper = "", "", ""
    if re.search(r'ALUM', up):
        family = "Aluminum"
    elif re.search(r'TITAN|TI-?6AL', up):
        family = "Titanium"
    elif re.search(r'STEEL|STAINLESS|\bSS\b|CRES', up) or STEEL_RE.search(body):
        family = "Steel"
    elif re.search(r'BRASS|COPPER|BRONZE', up):
        family = "Copper alloy"

    if family == "Aluminum":
        m = AL_RE.search(body)
        if m:
            alloy, temper = m.group(1), (m.group(2) or "")
    elif family == "Titanium":
        m = TI_RE.search(body)
        alloy = m.group(0).upper() if m else ""
    elif family == "Steel":
        m = STEEL_RE.search(body)
        alloy = m.group(1).upper() if m else ""
    if not temper:
        m = TEMPER_RE.search(body)
        temper = m.group(1).upper() if m else ""

    m = FORM_RE.search(body)
    form = m.group(1).upper() if m else ""
    dims = " ".join(f"{k.upper()}={v.rstrip('.')}" for k, v in DIM_RE.findall(body))
    if dims:
        u = UNIT_RE.search(body)
        if u:
            dims += " " + u.group(0).upper()
    c = CATNO_RE.search(extra)
    catno, qty, unit = (c.group(1), c.group(2), c.group(3) or "") if c else ("", "", "")

    if not (alloy or stds or form):
        return None
    return {"family": family, "alloy": alloy, "temper": temper, "form": form,
            "stds": "; ".join(stds), "dims": dims, "catno": catno, "qty": qty,
            "unit": unit, "src": " | ".join(r["text"] for r in mrows)}


# ---------- process extraction: heat treat / coating / paint / masking ----------
HT_GENERAL = "Heat treat (general)"
HT_TYPES = [
    ("Stress relief",    r"STRESS[- ]RELIEF|STRESS RELIEV"),
    ("Annealing",        r"ANNEAL"),
    ("Solution + aging", r"SOLUTION TREAT|\bAGING\b|AGE[- ]HARDEN|\bAGED\b"),
    ("Normalizing",      r"NORMALI[SZ]"),
    ("Carburizing",      r"CARBURI[SZ]"),
    ("Nitriding",        r"NITRID"),
    ("Hardening",        r"HARDEN|QUENCH"),
    ("Tempering",        r"\bTEMPER(?:ING|ED)?\b"),
    (HT_GENERAL,         r"HEAT[- ]?TREAT|\bHR[CB]\b|AMS[- ]?2759|H-6875"),
]

# each coating TYPE is its own record. Spec numbers are included so a bare spec also identifies the type.
COAT_TYPES = [
    ("Anodize",
     r"ANODI[SZ]\w*|MIL-A-8625|MIL-A-63576|AMS[- ]?24(?:68|69|7[0-2])"),
    ("Conversion coating (chem film)",
     r"CONVERSION COAT\w*|CHEM(?:ICAL)?[- ]FILM|ALODINE|IRIDITE|\bCHROMATE\b|MIL-(?:DTL|C)-5541|AMS[- ]?2473"),
    ("Electroless nickel",
     r"ELECTROLESS[- ]NI(?:CKEL)?\b|\bNI-?P\b|AMS[- ]?2404|(?:MIL|AMS)-C-26074|ASTM[- ]?B733"),
    ("Nickel plating (electrolytic/unspecified)",
     r"(?<!ELECTROLESS )(?<!ELECTROLESS-)(?<!ZINC-)(?<!ZINC )NICKEL(?!\s+ALLOY)|QQ-N-290|AMS[- ]?2403"),
    ("Zinc-nickel",     r"ZINC[- ]NICKEL|\bZN-?NI\b|ASTM[- ]?B841"),
    ("Zinc plating",    r"\bZINC\b(?![- ]NICKEL)|ASTM[- ]?B633|QQ-Z-325|AMS[- ]?2402"),
    ("Cadmium plating", r"CADMIUM|QQ-P-416|ASTM[- ]?B766|AMS[- ]?2400"),
    ("Chrome plating",  r"HARD CHROM\w*|CHROME PLAT\w*|CHROMIUM PLAT\w*|\bCHROME\b|AMS[- ]?2460|QQ-C-320"),
    ("Tin plating",     r"\bTIN PLAT\w*|MIL-T-10727|ASTM[- ]?B545"),
    ("Silver plating",  r"SILVER PLAT\w*|QQ-S-365|AMS[- ]?2410"),
    ("Gold plating",    r"GOLD PLAT\w*|MIL-G-45204|ASTM[- ]?B488"),
    ("Copper plating",  r"COPPER PLAT\w*|AMS[- ]?2418"),
    ("Passivation",     r"PASSIVAT\w*|AMS[- ]?2700|ASTM[- ]?A967|QQ-P-35"),
    ("Black oxide",     r"BLACK OXIDE|MIL-(?:DTL|C)-13924"),
    ("Phosphate",       r"PHOSPHAT\w*|MIL-(?:DTL|P)-16232|TT-C-490"),
    ("Dry film / PTFE", r"PTFE|TEFLON|DRY FILM|SOLID FILM LUBRICANT|MIL-(?:PRF|L)-(?:46010|8937)"),
    ("Plating (unspecified metal)", r"PLATING|PLATED|ELECTROPLAT\w*"),
    ("Electroless (other)",         r"ELECTROLESS"),
]
PLATING_SPECIFIC = {"Electroless nickel", "Nickel plating (electrolytic/unspecified)", "Zinc-nickel",
                    "Zinc plating", "Cadmium plating", "Chrome plating", "Tin plating",
                    "Silver plating", "Gold plating", "Copper plating"}
SUPPRESS = {"Electroless nickel": ["Electroless (other)"], "Zinc-nickel": ["Zinc plating"]}

PAINT_TYPES = [
    ("Powder coat",      r"POWDER[- ]COAT"),
    ("Primer",           r"PRIMER"),
    ("Top coat / paint", r"TOP ?COAT|FINISH COAT|ENAMEL|LACQUER|POLYURETHANE|\bPAINT"),
]
STRONG_RE = re.compile(r'ANODI|COATING|COATED|PLATING|PLATED|PAINT|PRIMER|HEAT[- ]?TREAT|PASSIVAT|'
                       r'MASK|CONVERSION|CHEM FILM|ALODINE', re.I)
SEQ_RE = re.compile(r'\b(?:AFTER|BEFORE|PRIOR TO|FOLLOWING)\s+(?:THE\s+)?'
                    r'(?:HEAT[- ]?TREAT\w*|HARDEN\w*|ANODI[SZ]\w*|PLAT\w*|PAINT\w*|COAT\w*|'
                    r'SURFACE TREATMENT|FINISH\w*)', re.I)

# masking
MASK_RE = re.compile(
    r'\bMASK(?:ED|ING)?\b'
    r'|FREE (?:OF|FROM) (?:PAINT|COAT\w*|PLAT\w*|ANODI\w*)'
    r'|DO NOT (?:COAT|PLATE|PAINT|ANODI[SZ]E)'
    r'|NO (?:COATING|PLATING|PAINT|ANODI[SZ]\w*) (?:IN|ON|AT|ALLOWED)'
    r'|NOT TO BE (?:COATED|PLATED|PAINTED|ANODI[SZ]ED)'
    r'|(?:EXCEPT|EXCLUDING|EXCLUDE)\s+(?:FOR\s+)?(?:THE\s+)?(?:THREAD\w*|HOLE\w*|BORE\w*|AREA\w*|SURFACE\w*)'
    r'|\bUNCOATED\b|\bUNPLATED\b|\bUNANODI[SZ]ED\b|\bUNPAINTED\b', re.I)
MASK_METHODS = [
    ("Tape",             r"\bTAPE\b"),
    ("Plug",             r"\bPLUG\w*"),
    ("Cap / boot",       r"\bCAPS?\b|\bCAPPED\b|\bBOOTS?\b"),
    ("Lacquer / stop-off", r"LACQUER|STOP[- ]?OFF"),
    ("Wax",              r"\bWAX\b"),
]
MASK_AREAS = [
    ("Threads",           r"THREAD"),
    ("Holes",             r"\bHOLES?\b"),
    ("Bores",             r"\bBORES?\b"),
    ("Internal surfaces", r"INTERNAL (?:SURFACE|AREA)"),
    ("Contact surfaces",  r"CONTACT (?:SURFACE|AREA)"),
    ("Grounding / electrical", r"GROUND|ELECTRICAL"),
    ("Sealing surfaces",  r"SEAL(?:ING)? SURFACE"),
    ("Datum",             r"\bDATUM"),
    ("Area per drawing",  r"AREAS?\s+(?:MARKED|SHOWN|INDICATED)|AS SHOWN|PER (?:VIEW|DETAIL|SECTION)"),
]
MASK_TARGET = [
    ("Anodize",            r"ANODI[SZ]"),
    ("Conversion coating", r"CONVERSION|CHEM(?:ICAL)?[- ]FILM|ALODINE"),
    ("Plating",            r"\bPLAT(?:E|ED|ING)\b|NICKEL"),
    ("Paint",              r"PAINT|PRIM(?:ER|ING)|POWDER"),
    ("Coating",            r"\bCOAT(?:ED|ING)?\b"),
]

LOC_RE = re.compile(r'\b(?:PER|SEE|IN)\s+(?:VIEW|DETAIL|SECTION|ZONE|SHEET)\b|\bAS SHOWN\b|\bAREA\s+[A-Z0-9]\b', re.I)
ASM_RE = re.compile(r'\bINSTALL|\bINSERT', re.I)
LAYER_WORDS = re.compile(r'FOLLOWED BY|\bTHEN\b|\bOVER\b|UNDER-?PLAT|\bSTRIKE\b', re.I)
TOTAL_PAREN = re.compile(r'\([^)]*\bTOTAL\b[^)]*\)', re.I)
LETTER = re.compile(r'^([A-Z])\s*[\.\)]\s+(.*)', re.S)

TYPE_RE = re.compile(r'\bTYPE\s*([IVX]+[ABC]?|\d+[A-Z]?)\b', re.I)
CLASS_RE = re.compile(r'\bCLASS\s*(\d+[A-Z]?|[A-Z](?![A-Z]))', re.I)
THK_RE = re.compile(
    r'(\d+(?:\.\d+)?(?:\s*(?:-|–|TO)\s*\d+(?:\.\d+)?)?)\s*'
    r'(MICRONS?|MICROMETERS?|[µμ]M|MILS?(?![-\w])|MM\b|INCH(?:ES)?\b)'
    r'(?:\s+(MIN(?:IMUM)?|MAX(?:IMUM)?)\b)?', re.I)
COATS_RE = re.compile(r'\b(ONE|TWO|THREE|[1-3])\s+COATS?\b', re.I)
COLOR_NUM = re.compile(r'\bCOLOU?R\s*(?:NO\.?\s*)?#?\s*(\d{4,5}[A-Z]?)', re.I)
COLOR_WORD = re.compile(r'\b(BLACK|WHITE|CLEAR|NATURAL|GREEN|GREY|GRAY|BLUE|RED|OLIVE DRAB|TAN|'
                        r'YELLOW|GOLD|ORANGE|SILVER)\b', re.I)
HARD_RE = re.compile(r'(\d{2}(?:\s*[-–]\s*\d{2})?)\s*(HR[CB])\b|\b(HR[CB])\s*[:=]?\s*(\d{2}(?:\s*[-–]\s*\d{2})?)', re.I)
CASE_RE = re.compile(r'CASE\s+DEPTH\s*[:=]?\s*([\d.]+(?:\s*[-–]\s*[\d.]+)?\s*(?:MM|INCH(?:ES)?|MICRONS?)?)', re.I)
TEMP_RE = re.compile(r'(?<![\w.\-])(?:(NOT EXCEED|MAX(?:IMUM)?|MIN(?:IMUM)?|AT|UP TO)\s+)?'
                     r'(\d{2,4}(?:\s*[-–]\s*\d{2,4})?)\s*(?:°|º|DEG\.?|DEGREES)?\s*([CF])\b', re.I)
NUMWORD = {"ONE": "1", "TWO": "2", "THREE": "3"}

# meanings from general knowledge of the specs - verify against the revisions you use
ANOD_TYPES = {"I": "chromic", "IB": "low-voltage chromic", "IC": "non-chromic",
              "II": "sulfuric", "IIB": "thin sulfuric", "III": "hard anodize"}
CONV_TYPES = {"I": "hexavalent chromium", "II": "non-hexavalent"}
CONV_CLASSES = {"1A": "max corrosion protection", "3": "electrical contact"}


def type_class(t):
    v = ["Type " + m.group(1).upper() for m in TYPE_RE.finditer(t)]
    v += ["Class " + m.group(1).upper() for m in CLASS_RE.finditer(t)]
    return "; ".join(uniq(v))


def thickness(t):
    out = []
    for m in THK_RE.finditer(t):
        val = re.sub(r'\s*(?:-|–|TO)\s*', '-', m.group(1), flags=re.I)
        val = re.sub(r'\s+', '', val)
        u = m.group(2).upper()
        if u.startswith("MICRO") or u[0] in "µμ":
            u = "µm"
        elif u.startswith("MIL"):
            u = "mil"
        else:
            u = u.lower()
        s = val + " " + u
        if m.group(3):
            s += " " + m.group(3)[:3].lower()
        if "TOTAL" in t[max(0, m.start() - 40):m.start()].upper():
            s += " (total)"
        out.append(s)
    return "; ".join(uniq(out))


def coats(t):
    m = COATS_RE.search(t)
    if not m:
        return ""
    g = m.group(1).upper()
    return NUMWORD.get(g, g)


def colors(t):
    out = ["#" + m.group(1).upper() for m in COLOR_NUM.finditer(t)]
    out += [m.group(1).title() for m in COLOR_WORD.finditer(t)]
    return "; ".join(uniq(out))


def hardness(t):
    out = []
    for m in HARD_RE.finditer(t):
        val = m.group(1) or m.group(4)
        sc = m.group(2) or m.group(3)
        out.append(re.sub(r'\s+', '', val) + " " + sc.upper())
    return "; ".join(uniq(out))


def temperature(t):
    out = []
    for m in TEMP_RE.finditer(t):
        pre = (m.group(1) or "").lower()
        val = re.sub(r'\s+', '', m.group(2))
        s = val + " " + m.group(3).upper()
        if pre:
            s = pre + " " + s
        out.append(s)
    return "; ".join(uniq(out))


def case_depth(t):
    m = CASE_RE.search(t)
    return re.sub(r'\s+', ' ', m.group(1)).strip() if m else ""


def find_coats(t):
    """Returns [(position, coating type)] sorted by position in the text."""
    hits = {}
    for name, rx in COAT_TYPES:
        m = re.search(rx, t, re.I)
        if m:
            hits[name] = m.start()
    for n in list(hits):
        for s in SUPPRESS.get(n, ()):
            hits.pop(s, None)
    if PLATING_SPECIFIC & set(hits):
        hits.pop("Plating (unspecified metal)", None)
    return sorted((p, n) for n, p in hits.items())


def coat_detail(name, w):
    up = w.upper()
    d = []
    if name == "Anodize":
        if re.search(r'HARD[- ]?(?:COAT|ANODI)', up):
            d.append("Hard anodize")
        if re.search(r'SULF', up):
            d.append("Sulfuric")
        if re.search(r'CHROMIC', up):
            d.append("Chromic")
        if "8625" in up:
            for m in TYPE_RE.finditer(w):
                tv = m.group(1).upper()
                if tv in ANOD_TYPES:
                    d.append(f"Type {tv} = {ANOD_TYPES[tv]}")
        if re.search(r'\bDY(?:E|ED)\b', up):
            d.append("Dyed")
        if re.search(r'\bSEAL', up):
            d.append("Sealed")
        if re.search(r'PTFE|TEFLON', up):
            d.append("PTFE-impregnated")
    elif name.startswith("Conversion"):
        if "5541" in up:
            for m in TYPE_RE.finditer(w):
                tv = m.group(1).upper()
                if tv in CONV_TYPES:
                    d.append(f"Type {tv} = {CONV_TYPES[tv]}")
            for m in CLASS_RE.finditer(w):
                cv = m.group(1).upper()
                if cv in CONV_CLASSES:
                    d.append(f"Class {cv} = {CONV_CLASSES[cv]}")
        if re.search(r'TRIVALENT', up):
            d.append("Trivalent chromium")
        if re.search(r'HEXAVALENT', up):
            d.append("Hexavalent chromium")
        if re.search(r'ROHS', up):
            d.append("RoHS")
    elif name == "Electroless nickel":
        m = re.search(r'(HIGH|MEDIUM|MID|LOW)[- ]PHOS', up)
        if m:
            d.append(m.group(1).title() + " phosphorus")
        if re.search(r'\bBAK(?:E|ING)\b|EMBRITTLEMENT', up):
            d.append("Bake mentioned")
    return "; ".join(d)


def mask_record(no, seg, common):
    methods = [n for n, rx in MASK_METHODS if re.search(rx, seg, re.I)]
    areas = [n for n, rx in MASK_AREAS if re.search(rx, seg, re.I)]
    before = [n for n, rx in MASK_TARGET if re.search(rx, seg, re.I)]
    if len(before) > 1 and "Coating" in before:
        before.remove("Coating")
    return {"no": no, "method": ", ".join(methods) or "Masking (method not stated)",
            "areas": ", ".join(areas), "before": " | ".join(before),
            "stds": "; ".join(get_standards(seg)), "flags": "; ".join(common), "src": seg}


def analyze_segment(no, seg):
    seqs = [m.group(0).upper() for m in SEQ_RE.finditer(seg)]
    t = SEQ_RE.sub(" ", seg)                    # "AFTER HEAT TREAT" is a sequence, not a process
    masked = bool(MASK_RE.search(t))
    t = MASK_RE.sub(" ", t)                     # "FREE OF PAINT" is masking, not paint
    tt = TOTAL_PAREN.sub(" ", t)

    common = ["SEQUENCE: " + s for s in seqs]
    if ASM_RE.search(seg):
        common.append("ASSEMBLY CONTEXT")
    if LOC_RE.search(seg):
        common.append("LOCATION IN DRAWING - REVIEW")

    recs, masks = [], []

    def add(process, subtype, w=None, extra_flags=(), **kw):
        w = t if w is None else w
        s = get_standards(w)
        fl = list(common) + list(extra_flags)
        if not s:
            fl.append("NO STANDARD")
        d = {"no": no, "process": process, "subtype": subtype, "stds": "; ".join(s),
             "layer": "", "detail": "", "typecls": "", "thick": "", "coats": "", "color": "",
             "hard": "", "temp": "", "case": "", "flags": "; ".join(fl), "src": seg}
        d.update(kw)
        recs.append(d)

    # --- heat treat
    ht = [n for n, rx in HT_TYPES if re.search(rx, tt, re.I)]
    if len(ht) > 1 and HT_GENERAL in ht:
        ht.remove(HT_GENERAL)
    if ht:
        add("Heat treat", " + ".join(ht), w=tt, hard=hardness(tt), temp=temperature(tt),
            case=case_depth(tt))

    # --- coatings: one record per coating type
    paint = [n for n, rx in PAINT_TYPES if re.search(rx, tt, re.I)]
    hits = find_coats(tt)
    if not hits and not paint and re.search(r'COATING|COATED', tt, re.I):
        hits = [(0, "Other / unspecified coating")]
    n = len(hits)
    for i, (s, name) in enumerate(hits):
        a = 0 if i == 0 else s
        b = hits[i + 1][0] if i < n - 1 else len(tt)
        w = tt[a:b]
        extra = []
        if n > 1:
            extra.append("MULTIPLE COATINGS - VERIFY FIELD ASSIGNMENT")
            if LAYER_WORDS.search(tt):
                extra.append("LAYER ORDER STATED IN TEXT")
        add("Coating", name, w=w, extra_flags=extra, layer=(f"{i + 1} of {n}" if n > 1 else ""),
            detail=coat_detail(name, w), typecls=type_class(w), thick=thickness(w),
            color=colors(w))

    # --- paint
    if paint:
        add("Paint", " + ".join(paint), typecls=type_class(t), thick=thickness(t),
            coats=coats(t), color=colors(t), temp=temperature(tt))

    # --- masking: separate process, happens before coating / paint
    if masked:
        masks.append(mask_record(no, seg, common))
    return recs, masks


def extract_processes(rows):
    out, masks = [], []
    for r in rows:
        text = r["text"]
        if "Material" in r["cat"] and not STRONG_RE.search(text):
            continue                              # plain material rows are not processes
        segs = [s.strip() for s in text.split("\n") if s.strip()]
        if not segs:
            continue
        items = []
        if len(segs) > 1:                         # header line + lettered sub-items (A., B., C.)
            if get_standards(segs[0]) or THK_RE.search(segs[0]):
                items.append((r["no"], segs[0]))
            for s in segs[1:]:
                m = LETTER.match(s)
                items.append((r["no"] + (m.group(1) if m else ""), m.group(2) if m else s))
        else:
            items.append((r["no"], segs[0]))
        for no, seg in items:
            recs, ms = analyze_segment(no, seg)
            out.extend(recs)
            masks.extend(ms)
    page_names = uniq([p["subtype"] for p in out if p["process"] in ("Coating", "Paint")])
    for m in masks:
        if not m["before"]:
            if page_names:
                m["before"] = " | ".join(page_names) + " (page-level)"
            else:
                m["flags"] = (m["flags"] + "; " if m["flags"] else "") + "NO COATING FOUND ON PAGE"
    return out, masks


def summarize(procs, masks):
    def j(items):
        return " | ".join(uniq([i for i in items if i]))

    def lab(r):
        s = r["subtype"]
        if r["typecls"]:
            s += " " + r["typecls"]
        if r["stds"]:
            s += " [" + r["stds"] + "]"
        return s

    ht = j([(r["subtype"] + " " + r["hard"]).strip() for r in procs if r["process"] == "Heat treat"])
    mk = j([(m["method"] + (": " + m["areas"] if m["areas"] else "")) for m in masks])
    co_recs = [r for r in procs if r["process"] == "Coating"]
    co = " → ".join(uniq([lab(r) for r in co_recs]))
    pa = j([lab(r) for r in procs if r["process"] == "Paint"])
    flow = (["Masking"] if masks else []) + uniq([r["subtype"] for r in co_recs]) \
        + (["Paint"] if pa else [])
    fl = sum(1 for r in procs + masks if r["flags"])
    return ht, mk, co, pa, " → ".join(flow), fl


# ---------- page reading ----------
def get_lines(page):
    out = []
    for b in page.get_text("dict")["blocks"]:
        if b["type"] != 0:
            continue
        for l in b["lines"]:
            t = clean("".join(s["text"] for s in l["spans"])).strip()
            if t:
                r = pymupdf.Rect(l["bbox"])
                out.append({"t": t, "r": r, "yc": (r.y0 + r.y1) / 2})
    return out


# ---------- mode 1: numbered table (NO. | DESCRIPTION) ----------
def find_table(lines):
    for h_no in lines:
        if h_no["t"].upper().rstrip(".") != "NO":
            continue
        for h_d in lines:
            if (h_d["t"].upper() == "DESCRIPTION" and abs(h_d["yc"] - h_no["yc"]) < 3
                    and h_d["r"].x0 > h_no["r"].x0):
                return h_no, h_d
    return None


def parse_table(lines, h_no, h_d):
    no_x, d_x, y0 = h_no["r"].x0, h_d["r"].x0, h_no["yc"]
    nums = [l for l in lines
            if abs(l["r"].x0 - no_x) < 6 and l["yc"] > y0 + 3 and NUM.match(l["t"])]
    nums.sort(key=lambda l: l["yc"])
    if not nums:
        return [], None
    rows = {id(n): {"no": n["t"], "desc": [], "extra": [], "rect": pymupdf.Rect(n["r"])}
            for n in nums}
    for l in lines:
        if l["yc"] <= y0 + 3 or any(l is n for n in nums) or l["r"].x0 < d_x - 4:
            continue
        n = min(nums, key=lambda n: abs(n["yc"] - l["yc"]))
        if abs(n["yc"] - l["yc"]) > 12:
            continue
        row = rows[id(n)]
        (row["desc"] if l["r"].x0 < d_x + 250 else row["extra"]).append(l)
        row["rect"] |= l["r"]

    result, section, box = [], ("", ""), pymupdf.Rect(h_no["r"])
    for n in nums:
        row = rows[id(n)]
        desc = " ".join(l["t"] for l in sorted(row["desc"], key=lambda l: (l["yc"], l["r"].x0)))
        extra = " ".join(l["t"] for l in sorted(row["extra"], key=lambda l: (l["yc"], l["r"].x0)))
        box |= row["rect"]
        if "." not in row["no"]:
            section = (row["no"], desc)
            continue
        cats = classify(desc) or classify(section[1])
        result.append({"section": f"{section[0]} {section[1]}".strip(), "no": row["no"],
                       "text": desc, "extra": extra,
                       "cat": ", ".join(cats) or section[1].title()})
    return result, box


# ---------- mode 2: classic NOTES block ----------
def parse_notes_block(page, lines):
    pw = page.rect.width
    best = None
    for a in [l for l in lines if ANCHOR.match(l["t"])]:
        cand = [l for l in lines if l is not a and l["r"].y0 >= a["r"].y1 - 2
                and a["r"].x0 - 6 <= l["r"].x0 < a["r"].x0 + 0.35 * pw]
        cand.sort(key=lambda l: l["yc"])
        rows, cur = [], None
        for l in cand:
            if cur and abs(l["yc"] - cur["yc"]) < 3:
                cur["parts"].append(l)
            else:
                cur = {"yc": l["yc"], "parts": [l]}
                rows.append(cur)
        merged, prev = [], a["r"]
        for row in rows:
            parts = sorted(row["parts"], key=lambda l: l["r"].x0)
            r = pymupdf.Rect(parts[0]["r"])
            for p in parts[1:]:
                r |= p["r"]
            if r.y0 - prev.y1 > 2.5 * max(r.height, 6):
                break
            merged.append({"t": " ".join(p["t"] for p in parts), "r": r,
                           "x0": parts[0]["r"].x0})
            prev = r
        if best is None or len(merged) > len(best[1]):
            best = (a, merged)
    if not best or not best[1]:
        return [], None

    a, merged = best
    TOP = re.compile(r'^(\d{1,2})\s*[\.\)]\s+(\S.*)')
    SUB = re.compile(r'^(\d{1,2}\.\d{1,2})\s+(\S.*)')
    LET = re.compile(r'^[A-Z]\s*[\.\)]\s+\S')
    items, last, top = [], 0, None
    box = pymupdf.Rect(a["r"])
    for m in merged:
        box |= m["r"]
        t = m["t"]
        left = abs(m["x0"] - a["r"].x0) <= 4
        mt, ms = TOP.match(t), SUB.match(t)
        if left and mt and int(mt.group(1)) > last:
            last = int(mt.group(1))
            top = {"section": "NOTES", "no": str(last), "text": mt.group(2), "parent": ""}
            items.append(top)
        elif left and ms and top:
            items.append({"section": f'{top["no"]} {top["text"][:30]}', "no": ms.group(1),
                          "text": ms.group(2), "parent": top["text"]})
        elif items:
            items[-1]["text"] += ("\n" if LET.match(t) else " ") + t
        else:
            items.append({"section": "NOTES", "no": "", "text": t, "parent": ""})

    out = []
    for i, it in enumerate(items):
        nxt = items[i + 1] if i + 1 < len(items) else None
        if nxt and nxt["no"].startswith(it["no"] + ".") and len(it["text"]) <= 40:
            continue
        it["cat"] = ", ".join(classify(it["text"]) or classify(it["parent"]))
        it["extra"] = ""
        out.append(it)
    return out, box


def process_page(page):
    lines = get_lines(page)
    if not lines:
        return "NEEDS_OCR", "", [], None
    t = find_table(lines)
    if t:
        rows, box = parse_table(lines, *t)
        if rows:
            return "OK", "ROUTING TABLE", rows, box
    rows, box = parse_notes_block(page, lines)
    if rows:
        return "OK", "NOTES BLOCK", rows, box
    return "NOT_FOUND", "", [], None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path", help="PDF file or folder")
    ap.add_argument("-o", "--out", default="notes.xlsx")
    ap.add_argument("--debug", action="store_true", help="save PNGs with the detected region")
    a = ap.parse_args()

    p = Path(a.path)
    files = [p] if p.is_file() else sorted(p.glob("*.pdf"))
    dbg = Path(a.out).resolve().parent / "debug"
    if a.debug:
        dbg.mkdir(exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"
    ws.append(["File", "Page", "Source", "Section", "No.", "Text", "Extra (cat no / qty)", "Category"])
    ws2 = wb.create_sheet("Summary")
    ws2.append(["File", "Page", "Status", "Source", "Items"])
    ws3 = wb.create_sheet("Material")
    ws3.append(["File", "Page", "Family", "Alloy", "Temper", "Form", "Standards",
                "Raw dimensions", "Cat no", "Qty", "Unit", "Source text"])
    ws4 = wb.create_sheet("Processes")
    ws4.append(["File", "Page", "No.", "Process", "Subtype", "Layer", "Detail",
                "Standards / specs", "Type / Class", "Thickness", "Coats", "Color",
                "Hardness", "Temperature", "Case depth", "Flags", "Source text"])
    ws6 = wb.create_sheet("Masking")
    ws6.append(["File", "Page", "No.", "Method", "Areas masked", "Before process",
                "Standards", "Flags", "Source text"])
    ws5 = wb.create_sheet("Process summary")
    ws5.append(["File", "Page", "Heat treat", "Masking", "Coatings (in order)", "Paint",
                "Flow", "Flagged"])

    for f in files:
        try:
            doc = pymupdf.open(f)
            for pno, page in enumerate(doc, 1):
                status, src, rows, box = process_page(page)
                for r in rows:
                    ws.append([clean(x) for x in
                               [f.name, pno, src, r["section"], r["no"],
                                r["text"], r["extra"], r["cat"]]])
                ws2.append([clean(f.name), pno, status, src, len(rows)])
                m = extract_material(rows)
                if m:
                    ws3.append([clean(x) for x in
                                [f.name, pno, m["family"], m["alloy"], m["temper"], m["form"],
                                 m["stds"], m["dims"], m["catno"], m["qty"], m["unit"], m["src"]]])
                procs, masks = extract_processes(rows)
                for q in procs:
                    ws4.append([clean(x) for x in
                                [f.name, pno, q["no"], q["process"], q["subtype"], q["layer"],
                                 q["detail"], q["stds"], q["typecls"], q["thick"], q["coats"],
                                 q["color"], q["hard"], q["temp"], q["case"], q["flags"], q["src"]]])
                for mk in masks:
                    ws6.append([clean(x) for x in
                                [f.name, pno, mk["no"], mk["method"], mk["areas"], mk["before"],
                                 mk["stds"], mk["flags"], mk["src"]]])
                if rows:
                    ht, mks, co, pa, flow, fl = summarize(procs, masks)
                    ws5.append([clean(x) for x in [f.name, pno, ht, mks, co, pa, flow, fl]])
                if a.debug:
                    if box is not None:
                        page.draw_rect(box, color=(1, 0, 0), width=1.5)
                    page.insert_text((20, 20), status, fontsize=14, color=(1, 0, 0))
                    page.get_pixmap(dpi=100).save(dbg / f"{f.stem}_p{pno}.png")
        except Exception as e:
            ws2.append([clean(f.name), "", clean(f"ERROR: {e}"), "", 0])

    for w, widths in ((ws, [24, 6, 15, 24, 7, 90, 26, 18]),
                      (ws2, [24, 6, 14, 15, 8]),
                      (ws3, [24, 6, 12, 14, 9, 11, 34, 24, 13, 8, 6, 70]),
                      (ws4, [24, 6, 7, 11, 30, 8, 34, 34, 18, 26, 6, 14, 12, 16, 12, 34, 80]),
                      (ws6, [24, 6, 7, 22, 34, 40, 24, 30, 80]),
                      (ws5, [24, 6, 30, 40, 50, 44, 50, 9])):
        for i, wd in enumerate(widths):
            w.column_dimensions[chr(65 + i)].width = wd
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "A2"
    for sheet, col in ((ws, 6), (ws4, 17), (ws6, 9)):
        for row in sheet.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(a.out)
    print(f"{len(files)} file(s) -> {a.out}")


if __name__ == "__main__":
    main()